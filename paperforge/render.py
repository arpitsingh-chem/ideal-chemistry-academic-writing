"""Render a composed Document to Markdown, LaTeX or Word."""
from __future__ import annotations

import re
from pathlib import Path

from .draft import CITE_MARK, Document
from .project import Project, ProjectError


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def to_markdown(doc: Document) -> str:
    lines = [f"# {doc.title}", ""]
    affs = _affiliation_lines(doc)
    author_line = ", ".join(
        a["name"] + (f" ({a['affiliation']})" if a.get("affiliation") else "") + ("*" if a.get("corresponding") else "")
        for a in doc.authors
    )
    if author_line:
        lines += [author_line, ""]
    for key, text in affs.items():
        lines.append(f"^{key}^ {text}")
    corr = [a for a in doc.authors if a.get("corresponding")]
    if corr:
        lines.append("")
        for a in corr:
            if a.get("email"):
                lines.append(f"*Corresponding author: {a['email']}" + (f" (ORCID {a['orcid']})" if a.get("orcid") else ""))
    lines += ["", "## Abstract", "", doc.abstract, ""]
    if doc.keywords:
        lines += ["**Keywords:** " + "; ".join(doc.keywords), ""]
    if doc.profile.toc_width_cm:
        lines += [f"> TOC graphic target: {doc.profile.toc_width_cm} × {doc.profile.toc_height_cm} cm"
                  + (f"; TOC text ≤{doc.profile.toc_text_max_chars} characters" if doc.profile.toc_text_max_chars else ""),
                  ""]
    num = {r["key"]: r["n"] for r in doc.references}

    for sec in doc.sections:
        lines += [f"## {sec.heading}", ""]
        for p in sec.paragraphs:
            if p["text"].startswith("__SUB__"):
                lines += [f"### {p['text'][7:]}", ""]
                continue
            txt = CITE_MARK.sub(lambda m: "[" + ",".join(str(num.get(k, k)) for k in m.group(1).split(",")) + "]", p["text"])
            lines += [_md_cites(txt, p.get("cites") or [], num), ""]
        for tid in sec.tables:
            t = doc.tables.get(tid)
            if not t:
                continue
            lines += [f"**Table {_table_number(doc, tid)}.** {t.caption}", ""]
            lines.append("| " + " | ".join(human_header(h) for h in t.headers) + " |")
            lines.append("|" + "|".join("---" for _ in t.headers) + "|")
            for row in t.rows:
                lines.append("| " + " | ".join(row) + " |")
            if t.footnote:
                lines += ["", f"*{t.footnote}*"]
            lines.append("")

    if doc.figures:
        lines += ["## Figures", ""]
        for f in doc.figures:
            lines += [f"![{f['caption']}]({f['file']})", "", f"**{f['label']}.** {f['caption']}", ""]

    if doc.toc_text:
        lines += ["## Table of Contents Entry", "", doc.toc_text, ""]

    lines += ["## References", ""]
    for r in doc.references:
        lines.append(f"{r['n']}. {r['text']}")
    lines.append("")

    bm = doc.back_matter or {}
    if bm.get("acknowledgements") or bm.get("funding"):
        lines += ["## Acknowledgements", "", bm.get("acknowledgements") or bm.get("funding", ""), ""]
    for title, key in (("Author Contributions", "credit"), ("Conflict of Interest", "coi"),
                       ("Data Availability", "data_availability"), ("Safety Statement", "safety"),
                       ("Ethics Approval", "ethics_approval"), ("AI-Use Disclosure", "ai_disclosure")):
        if bm.get(key):
            lines += [f"## {title}", "", str(bm[key]), ""]
    return "\n".join(lines)


def _md_cites(text: str, cites: list, num: dict) -> str:
    if not cites:
        return text
    refs = [num.get(c.strip(), c.strip()) for c in cites if str(c).strip()]
    if not refs:
        return text
    marker = "[" + ",".join(str(r) for r in refs) + "]"
    # attach the marker to the previous word
    return text.rstrip() + marker


def human_header(name: str) -> str:
    """yield_percent -> Yield (%); conversion_percent -> Conversion (%); aryl_group -> Aryl group."""
    h = str(name).strip().replace("_", " ")
    if h.endswith(" percent"):
        h = h[: -len(" percent")] + " (%)"
    elif h.endswith(" (percent)"):
        h = h.replace(" (percent)", " (%)")
    return h[:1].upper() + h[1:] if h else h


def caption_tex(text: str) -> str:
    """Escape a caption, converting a trailing '^a' marker into a superscript."""
    s = str(text)
    s = re.sub(r"\^([a-z])\b", r"@@fn\1@@", s)
    out = tex_escape(s)
    return re.sub(r"@@fn([a-z])@@", r"\\textsuperscript{\1}", out)


def _table_number(doc: Document, tid: str) -> int:
    tids = list(doc.tables)
    return tids.index(tid) + 1 if tid in tids else 1


def _affiliation_lines(doc: Document) -> dict:
    aff = doc.affiliations or {}
    if isinstance(aff, list):
        return {i + 1: a for i, a in enumerate(aff)}
    return {str(k): str(v) for k, v in aff.items()}


# ---------------------------------------------------------------------------
# LaTeX
# ---------------------------------------------------------------------------

_GREEK = {
    "α": r"$\alpha$", "β": r"$\beta$", "γ": r"$\gamma$", "δ": r"$\delta$", "ε": r"$\varepsilon$",
    "λ": r"$\lambda$", "μ": r"$\mu$", "µ": r"$\mu$", "ν": r"$\nu$", "π": r"$\pi$", "σ": r"$\sigma$",
    "Φ": r"$\Phi$", "φ": r"$\phi$", "ω": r"$\omega$", "Ω": r"$\Omega$", "τ": r"$\tau$",
    "Å": r"\AA{}", "°": r"\textdegree{}", "±": r"$\pm$", "×": r"$\times$", "≈": r"$\approx$",
    "≤": r"$\leq$", "≥": r"$\geq$", "−": r"$-$", "→": r"$\rightarrow$",
    "¹": r"\textsuperscript{1}", "²": r"\textsuperscript{2}", "³": r"\textsuperscript{3}",
    "⁴": r"\textsuperscript{4}", "⁰": r"\textsuperscript{0}",
    "–": "--", "—": "---", "’": "'", "‘": "'", "“": "``", "”": "''", "…": r"\ldots{}",
    "≡": r"$\equiv$", "≫": r"$\gg$", "≪": r"$\ll$", "·": r"$\cdot$",
    "⁻": r"$^{-}$", "⁺": r"$^{+}$", "ν̃": r"$\tilde{\nu}$", "κ": r"$\kappa$",
    "Δ": r"$\Delta$", "η": r"$\eta$", "ρ": r"$\rho$", "θ": r"$\theta$",
}


def tex_escape(text: str) -> str:
    """Escape plain text for LaTeX while keeping common chemistry symbols.

    Chemistry tokens (¹H, ¹³C, cm⁻¹, ν̃, molecular formulae) are protected
    before escaping and restored as proper LaTeX/mhchem afterwards.
    """
    if not text:
        return ""
    s = str(text)
    protected: list[tuple[str, str]] = []

    def stash(token: str, latex: str) -> None:
        key = f"@@pf{len(protected)}@@"
        protected.append((key, latex))
        nonlocal s
        s = s.replace(token, key)

    stash("¹³C{¹H}", r"\ce{^{13}C\{^1H\}}")   # before the single-nucleus tokens
    stash("¹H", r"\ce{^1H}")
    stash("¹³C", r"\ce{^{13}C}")
    stash("cm⁻¹", r"cm$^{-1}$")
    s = re.sub(r"([A-Za-z])−1\b", r"\1@@supminus1@@", s)
    protected.append(("@@supminus1@@", r"$^{-1}$"))
    stash("ν̃", r"$\tilde{\nu}$")
    # molecular formulae such as C14H10N2O or C15H12N2O2
    for m in sorted(set(re.findall(r"\b(?:[A-Z][a-z]?\d*){2,}\b", s)), key=len, reverse=True):
        if any(c.isdigit() for c in m) and re.fullmatch(r"(?:[A-Z][a-z]?\d*)+", m) and len(m) >= 4:
            stash(m, r"\ce{" + m + "}")

    out = []
    for ch in s:
        if ch in _GREEK:
            out.append(_GREEK[ch])
        elif ch in "&%$#_{}":
            out.append("\\" + ch)
        elif ch == "\\":
            out.append(r"\textbackslash{}")
        elif ch == "~":
            out.append(r"\textasciitilde{}")
        elif ch == "\u0303":  # stray combining tilde
            continue
        else:
            out.append(ch)
    s = "".join(out)
    s = re.sub(r"(?<![\w$])\^(\d+)(?![\w{])", r"\\textsuperscript{\1}", s)
    for key, latex in protected:
        s = s.replace(key, latex)
    return s


def to_latex(doc: Document, project: Project | None = None, si: bool = False) -> str:
    p = doc.profile
    cls = "article"
    preamble = r"""\documentclass[12pt,a4paper]{article}
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage{mathptmx}
\usepackage[margin=2.5cm]{geometry}
\usepackage{setspace}
\onehalfspacing
\usepackage{graphicx}
\usepackage{booktabs,array,multirow}
\usepackage[version=4]{mhchem}
\usepackage{siunitx}
\usepackage[font=small,labelfont=bf]{caption}
\usepackage[numbers,sort&compress]{natbib}
\usepackage{lineno}
\usepackage{textcomp}
\usepackage[colorlinks=true,linkcolor=blue,citecolor=blue,urlcolor=blue]{hyperref}
\bibliographystyle{unsrtnat}
"""
    if si:
        return _si_latex(doc, project)
    affs = _affiliation_lines(doc)
    author_tex = []
    for a in doc.authors:
        sup = str(a.get("affiliation", ""))
        line = tex_escape(a["name"]) + (rf"\textsuperscript{{{sup}}}" if sup else "")
        if a.get("corresponding") and a.get("email"):
            orcid = f" (ORCID: {a['orcid']})" if a.get("orcid") else ""
            line += r"\thanks{*Corresponding author. E-mail: " + tex_escape(a["email"]) + tex_escape(orcid) + "}"
        author_tex.append(line)
    aff_lines = " \\\\ ".join(
        rf"\textsuperscript{{{k}}} {tex_escape(v)}" for k, v in affs.items()
    )
    aff_tex = r"\parbox{0.92\textwidth}{\centering " + aff_lines + "}"
    title = tex_escape(doc.title)
    keywords = "; ".join(doc.keywords)
    num = {r["key"]: r["n"] for r in doc.references}

    body: list[str] = []
    body.append(r"\begin{abstract}")
    body.append(tex_escape(doc.abstract))
    body.append(r"\end{abstract}")
    body.append("")
    if keywords:
        body.append(r"\noindent\textbf{Keywords:} " + tex_escape(keywords) + r"\\[6pt]")
    fig_counter = 0
    for sec in doc.sections:
        body.append(rf"\section{{{tex_escape(sec.heading)}}}")
        for para in sec.paragraphs:
            if para["text"].startswith("__SUB__"):
                body.append(rf"\subsection{{{tex_escape(para['text'][7:])}}}")
                continue
            txt = tex_escape(para["text"])
            # tex_escape has already run over the marker, so keys with "_" arrive
            # as "unicode\_title"; unescape them or \cite dies with \endcsname.
            txt = CITE_MARK.sub(lambda m: r"\cite{" + m.group(1).replace("\\_", "_") + "}", txt)
            cites = [c.strip() for c in (para.get("cites") or []) if str(c).strip()]
            if cites:
                txt = txt.rstrip() + r"~\cite{" + ",".join(cites) + "}"
            body.append(txt)
            body.append("")
        for tid in sec.tables:
            t = doc.tables.get(tid)
            if not t:
                continue
            body.append(r"\begin{table}[htbp]")
            body.append(r"\centering")
            body.append(rf"\caption{{{caption_tex(t.caption)}}}")
            body.append(r"\small")
            if len(t.headers) >= 5:
                body.append(r"\resizebox{\linewidth}{!}{%")
            body.append(r"\begin{tabular}{l" + "c" * (len(t.headers) - 1) + "}")
            body.append(r"\toprule")
            body.append(" & ".join(tex_escape(human_header(h)) for h in t.headers) + r" \\")
            body.append(r"\midrule")
            for row in t.rows:
                body.append(" & ".join(tex_escape(c) for c in row) + r" \\")
            body.append(r"\bottomrule")
            body.append(r"\end{tabular}")
            if len(t.headers) >= 5:
                body.append("}")
            if t.footnote:
                body.append(rf"\par\smallskip\footnotesize {tex_escape(t.footnote)}")
            body.append(r"\end{table}")
            body.append("")
    for f in doc.figures:
        fig_counter += 1
        body.append(r"\begin{figure}[htbp]")
        body.append(r"\centering")
        body.append(rf"\includegraphics[width=0.85\linewidth]{{{f['file']}}}")
        body.append(rf"\caption{{{tex_escape(f['caption'])}}}")
        body.append(rf"\label{{fig:{f['id']}}}")
        body.append(r"\end{figure}")
        body.append("")
    if doc.toc_text:
        body.append(r"\paragraph{Table of contents entry.} " + tex_escape(doc.toc_text))
    bm = doc.back_matter or {}
    body.append(r"\paragraph{Supporting Information.}")
    si_items = ["General experimental methods", "NMR spectra for all new compounds",
                "HRMS spectra", "References"]
    if project and any(c.crystal for c in project.compounds):
        si_items.insert(3, "crystallographic data (CIF)")
    body.append("Supporting Information: " + tex_escape(", ".join(si_items)) + " (PDF).")
    if bm.get("acknowledgements") or bm.get("funding"):
        body.append(r"\paragraph{Acknowledgements.} " + tex_escape(bm.get("acknowledgements") or bm.get("funding", "")))
    if bm.get("coi"):
        body.append(r"\paragraph{Conflict of interest.} " + tex_escape(bm["coi"]))
    if bm.get("data_availability"):
        body.append(r"\paragraph{Data availability.} " + tex_escape(bm["data_availability"]))
    if bm.get("credit"):
        body.append(r"\paragraph{Author contributions.} " + tex_escape(bm["credit"]))
    elif project:
        credits = []
        for a in project.authors:
            if a.credit:
                credits.append(f"{a.name}: {', '.join(a.credit)}.")
        if credits:
            body.append(r"\paragraph{Author contributions.} " + tex_escape(" ".join(credits)))
    if bm.get("ai_disclosure"):
        body.append(r"\paragraph{AI-use disclosure.} " + tex_escape(bm["ai_disclosure"]))

    author_block = ", ".join(author_tex) + r"\\[6pt]" + "\n" + aff_tex
    head = (preamble + "\n"
            + "\\title{" + title + "}\n"
            + "\\author{" + author_block + "}\n"
            + "\\date{}\n"
            + "\\begin{document}\n"
            + "\\maketitle\n"
            + "\\thispagestyle{empty}\n"
            + "\\linenumbers\n")
    tail = "\n\\bibliography{refs/refs}\n\\end{document}\n"
    return head + "\n\n".join(body) + tail


def _si_latex(doc: Document, project: Project | None) -> str:
    from .draft import characterisation_line

    head = r"""\documentclass[12pt,a4paper]{article}
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage{mathptmx}
\usepackage[margin=2.5cm]{geometry}
\usepackage{graphicx}
\usepackage{booktabs}
\usepackage[version=4]{mhchem}
\usepackage{siunitx}
\usepackage{textcomp}
\usepackage[colorlinks=true]{hyperref}
\title{Supporting Information\\[6pt]\large %s}
\author{%s}
\date{}
\begin{document}
\maketitle
""" % (tex_escape(doc.title), tex_escape(", ".join(a["name"] for a in doc.authors)))
    body = [r"\section*{Contents}", r"\begin{itemize}",
            r"\item General experimental methods",
            r"\item Synthetic procedures and characterisation",
            r"\item NMR spectra", r"\item HRMS spectra", r"\end{itemize}", r"\newpage"]
    if project:
        bm_methods = (project.narrative or {}).get("general_methods", "")
        body += [r"\section*{General experimental methods}",
                 tex_escape(bm_methods or "All reagents were used as received. NMR spectra were recorded on a Bruker Avance spectrometer.")]
        body.append(r"\section*{Synthetic procedures and characterisation}")
        for c in project.compounds:
            if c.known:
                continue
            body.append(rf"\subsection*{{Compound {tex_escape(c.id)}}}")
            proc = c.procedure or "Prepared according to the general procedure."
            body.append(tex_escape(proc))
            char = characterisation_line(c)
            if char:
                body.append(tex_escape(char))
        if any(c.crystal for c in project.compounds):
            body.append(r"\section*{Crystallographic data}")
            for c in project.compounds:
                if c.crystal and c.ccdc:
                    body.append(tex_escape(
                        f"CCDC {c.ccdc} contains the supplementary crystallographic data for compound {c.id}."))
    return head + "\n".join(body) + "\n\\end{document}\n"


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def to_docx(doc: Document, path: str | Path, project: Project | None = None) -> Path:
    try:
        from docx import Document as Docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover
        raise ProjectError("python-docx is required for .docx output (`pip install python-docx`)") from exc

    d = Docx()
    for section in d.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    t = d.add_paragraph()
    run = t.add_run(doc.title)
    run.bold = True
    run.font.size = Pt(16)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    authors = ", ".join(a["name"] + ("*" if a.get("corresponding") else "") for a in doc.authors)
    ap = d.add_paragraph(authors)
    ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, v in _affiliation_lines(doc).items():
        pp = d.add_paragraph(f"{k} {v}")
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.runs[0].font.size = Pt(10)
    for a in doc.authors:
        if a.get("corresponding") and a.get("email"):
            pp = d.add_paragraph(f"*E-mail: {a['email']}" + (f"  ORCID: {a['orcid']}" if a.get("orcid") else ""))
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pp.runs[0].italic = True

    d.add_heading("Abstract", level=1)
    d.add_paragraph(doc.abstract)
    if doc.keywords:
        p = d.add_paragraph()
        p.add_run("Keywords: ").bold = True
        p.add_run("; ".join(doc.keywords))

    num = {r["key"]: r["n"] for r in doc.references}
    for sec in doc.sections:
        d.add_heading(sec.heading, level=1)
        for para in sec.paragraphs:
            if para["text"].startswith("__SUB__"):
                d.add_heading(para["text"][7:], level=2)
                continue
            text = CITE_MARK.sub(lambda m: "[" + ",".join(str(num.get(k, k)) for k in m.group(1).split(",")) + "]", para["text"])
            cites = [c.strip() for c in (para.get("cites") or []) if str(c).strip()]
            if cites:
                text = text.rstrip() + "[" + ",".join(str(num.get(c, c)) for c in cites) + "]"
            d.add_paragraph(text)
        for tid in sec.tables:
            table = doc.tables.get(tid)
            if not table:
                continue
            cap = d.add_paragraph()
            cap.add_run(f"Table {list(doc.tables).index(tid)+1}. {table.caption}").bold = True
            tbl = d.add_table(rows=1, cols=len(table.headers))
            tbl.style = "Table Grid"
            for i, h in enumerate(table.headers):
                cell = tbl.rows[0].cells[i]
                cell.text = human_header(h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in table.rows:
                cells = tbl.add_row().cells
                for i, val in enumerate(row[: len(table.headers)]):
                    cells[i].text = val
            if table.footnote:
                fp = d.add_paragraph(table.footnote)
                fp.runs[0].font.size = Pt(9)
                fp.runs[0].italic = True

    for f in doc.figures:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img = project.path(f["file"]) if project else Path(f["file"])
        if img.exists() and img.suffix.lower() in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
            try:
                p.add_run().add_picture(str(img), width=Cm(14))
            except Exception:
                p.add_run(f"[figure: {f['file']}]")
        else:
            p.add_run(f"[figure: {f['file']}]")
        cap = d.add_paragraph()
        cap.add_run(f"{f['label']}. ").bold = True
        cap.add_run(f["caption"])

    if doc.references:
        d.add_heading("References", level=1)
        for r in doc.references:
            p = d.add_paragraph(f"{r['n']}. {r['text']}")
            p.paragraph_format.space_after = Pt(2)

    bm = doc.back_matter or {}
    for title, key in (("Acknowledgements", "acknowledgements"), ("Funding", "funding"),
                       ("Author Contributions", "credit"), ("Conflict of Interest", "coi"),
                       ("Data Availability", "data_availability"), ("Safety Statement", "safety"),
                       ("Ethics Approval", "ethics_approval"), ("AI-Use Disclosure", "ai_disclosure")):
        if bm.get(key):
            d.add_heading(title, level=2)
            d.add_paragraph(str(bm[key]))

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)
    return out


def to_plain(doc: Document) -> str:
    lines = [doc.title, "", ", ".join(a["name"] for a in doc.authors), "", "ABSTRACT", doc.abstract, ""]
    for sec in doc.sections:
        lines += [sec.heading.upper(), ""]
        for p in sec.paragraphs:
            lines.append(CITE_MARK.sub("", p["text"]).replace("__SUB__", ""))
            lines.append("")
    return "\n".join(lines)
